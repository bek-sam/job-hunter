"""Render a cover letter to PDF (LibreOffice) and DOCX (python-docx). 0 LLM tokens.
Files are named First_Last_Cover_Letter.pdf/.docx (applicant.file) -- never the company name."""
import html, os
import applicant as P
from docx import Document
from docx.shared import Pt, Inches

HEADER = (P.NAME, P.get('cover_letter_header', default=''))


def to_pdf(text, folder):
    """DOCX -> PDF with LibreOffice (same as gen_v5). Not Playwright: apply_prep calls this inside its own
    sync_playwright() block, where a second sync Playwright can't start."""
    import subprocess
    dp = to_docx(text, folder)
    prof = 'file:///tmp/lo_cl_' + str(os.getpid())
    subprocess.run(['/Applications/LibreOffice.app/Contents/MacOS/soffice', '-env:UserInstallation=' + prof, '--headless', '--convert-to', 'pdf', '--outdir', folder, dp],
                   capture_output=True, timeout=180)
    out = dp[:-5] + '.pdf'
    if not os.path.exists(out): raise RuntimeError('soffice did not produce the cover letter PDF')
    return out


def to_docx(text, folder):
    d = Document(); s = d.sections[0]
    s.left_margin = s.right_margin = Inches(0.9); s.top_margin = s.bottom_margin = Inches(0.8)
    st = d.styles['Normal']; st.font.name = 'Carlito'; st.font.size = Pt(11)
    r = d.add_paragraph().add_run(HEADER[0]); r.bold = True; r.font.size = Pt(18)
    d.add_paragraph(HEADER[1])
    for p in text.split('\n\n'):
        if p.strip(): d.add_paragraph(p.strip())
    d.core_properties.author = HEADER[0]; d.core_properties.title = f'{P.NAME} — Cover Letter'
    out = os.path.join(folder, P.file('Cover_Letter', 'docx')); d.save(out)
    return out
